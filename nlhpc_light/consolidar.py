# -*- coding: utf-8 -*-
"""
Junta los resultados de NLHPC (una carpeta por instancia) en tablas unicas y,
si existen ambas formulaciones, las compara par a par.

Uso (desde la raiz del repo):
    python nlhpc_light/consolidar.py                       # T3600 por defecto
    python nlhpc_light/consolidar.py --limite 60

Lee   corridas_nlhpc/light_v1/T<lim>/*/i*/resultados_familia.csv
      corridas_nlhpc/baseline_v1/T<lim>/*/i*/resultados_familia.csv
Crea  corridas_nlhpc/consolidado_T<lim>/
        resultados_light.csv, resultados_baseline.csv   (una fila por instancia)
        resumen_por_familia.txt                          (por formulacion)
        comparacion_pares.csv / .txt                     (si hay ambas)
        resumen_resultados.xlsx                          (Excel, ver abajo)

Excel resumen_resultados.xlsx (se regenera completo en cada llamada; cada tarea
de NLHPC lo llama al terminar, asi que se va llenando a medida que hay soluciones):
  Hoja "Resumen"  una fila por familia y formulacion:
      instancia (familia L*_P*_C*), formulacion, replicas terminadas, optimas,
      gap promedio (%), gap desv. estandar (%), tiempo promedio (s),
      tiempo desv. estandar (s)
  Hoja "Detalle"  una fila por instancia resuelta (estado, gap, tiempo, cotas...)
  Hoja "Asignaciones"  una fila por estudiante e instancia: curso nuevo y si
      quedo con al menos un preferido (tomado de cada soluciones/sol_*.txt)
Ademas copia todas las soluciones .txt a consolidado_T<lim>/soluciones/<formulacion>/
La desviacion estandar es muestral (n-1), igual que resumen_familia.csv.
Gap: el relativo de SCIP en %; las replicas sin solucion (gap infinito) no
entran en el promedio y se cuentan en la columna "sin incumbente".
"""
import argparse
import csv
import math
import statistics as st
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
CAMPOS = ["instancia", "L", "P", "C", "replica", "estado", "tiempo_scip_seg",
          "primal_bound", "dual_bound", "gap_pct", "nodos", "T", "suma_z",
          "no_satisfechos", "solucion_auditada", "variables", "restricciones"]


def f(x):
    try:
        v = float(x)
        return v if math.isfinite(v) else None
    except (TypeError, ValueError):
        return None


def carpeta(form, limite):
    return RAIZ / "corridas_nlhpc" / f"{form}_v1" / f"T{limite}"


def leer_asignaciones(form, limite):
    """[(instancia, fila de [ASIGNACION])...] desde soluciones/sol_*.txt."""
    out = []
    for p in sorted(carpeta(form, limite).glob("*/i*/soluciones/sol_*.txt")):
        inst = p.stem[len("sol_"):]
        lineas = p.read_text(encoding="utf-8").splitlines()
        if "[ASIGNACION]" not in lineas:
            continue
        k = lineas.index("[ASIGNACION]")
        cab = lineas[k + 1].split(";")
        for ln in lineas[k + 2:]:
            if not ln.strip() or ln.startswith("[") or ln.startswith("#"):
                break
            out.append((inst, dict(zip(cab, ln.split(";")))))
    return out


def copiar_soluciones(form, limite, destino):
    import shutil
    n = 0
    for p in carpeta(form, limite).glob("*/i*/soluciones/sol_*.txt"):
        d = destino / form
        d.mkdir(parents=True, exist_ok=True)
        shutil.copy2(p, d / p.name)
        n += 1
    return n


def leer(form, limite):
    base = RAIZ / "corridas_nlhpc" / f"{form}_v1" / f"T{limite}"
    filas = {}
    for p in sorted(base.glob("*/i*/resultados_familia.csv")):
        with p.open(encoding="utf-8-sig", newline="") as fh:
            for r in csv.DictReader(fh):
                filas[r["instancia"]] = {k: r.get(k, "") for k in CAMPOS}
    return filas


def familia(r):
    return f"L{r['L']}_P{r['P']}_C{r['C']}"


def media_geo(v, s=10.0):
    v = [x for x in v if x is not None]
    return math.exp(sum(math.log(x + s) for x in v) / len(v)) - s if v else None


def fmt(x, d=1):
    return "-" if x is None else f"{x:.{d}f}"


def escribir_csv(ruta, filas, campos):
    with ruta.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=campos)
        w.writeheader()
        w.writerows(filas)


def resumen(filas, titulo):
    lin = [titulo, "=" * len(titulo),
           f"{'familia':<14}{'n':>3}{'optimas':>9}{'t_medio_s':>11}{'t_geo_s':>9}"
           f"{'gap_medio_%':>12}{'auditadas':>11}"]
    fams = {}
    for r in filas.values():
        fams.setdefault(familia(r), []).append(r)
    for fam in sorted(fams, key=lambda s: [int(t[1:]) for t in s.split("_")]):
        rs = fams[fam]
        t = [f(r["tiempo_scip_seg"]) for r in rs]
        g = [f(r["gap_pct"]) for r in rs if f(r["gap_pct"]) is not None]
        opt = sum(r["estado"] == "optimal" for r in rs)
        aud = sum(r["solucion_auditada"] == "TRUE" for r in rs)
        lin.append(f"{fam:<14}{len(rs):>3}{opt:>9}{fmt(st.mean([x for x in t if x is not None]) if t else None):>11}"
                   f"{fmt(media_geo(t)):>9}{fmt(st.mean(g) if g else None, 2):>12}{aud:>11}")
    return lin


def estadisticas(valores):
    v = [x for x in valores if x is not None]
    if not v:
        return None, None
    return st.mean(v), (st.stdev(v) if len(v) > 1 else None)


def escribir_excel(ruta, datos, limite):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("AVISO: falta openpyxl (pip install openpyxl); no se genera el Excel.")
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    cab = ["instancia", "formulacion", "replicas terminadas", "optimas", "sin incumbente",
           "gap promedio (%)", "gap desv. estandar (%)",
           "tiempo promedio (s)", "tiempo desv. estandar (s)"]
    ws.append(cab)
    orden = lambda fam: [int(t[1:]) for t in fam.split("_")]  # noqa: E731
    for form in ("light", "baseline"):
        fams = {}
        for r in datos.get(form, {}).values():
            fams.setdefault(familia(r), []).append(r)
        for fam in sorted(fams, key=orden):
            rs = fams[fam]
            gp, gs = estadisticas([f(r["gap_pct"]) for r in rs])
            tp, ts = estadisticas([f(r["tiempo_scip_seg"]) for r in rs])
            ws.append([fam, form, len(rs), sum(r["estado"] == "optimal" for r in rs),
                       sum(f(r["gap_pct"]) is None for r in rs), gp, gs, tp, ts])
    dt = wb.create_sheet("Detalle")
    cab_d = ["instancia", "familia", "formulacion", "replica", "estado", "gap (%)",
             "tiempo SCIP (s)", "primal", "dual", "nodos", "T", "suma_z", "auditada"]
    dt.append(cab_d)
    for form in ("light", "baseline"):
        for n in sorted(datos.get(form, {}), key=lambda x: [int(t) if t.isdigit() else t
                                                            for t in x.split("_")]):
            r = datos[form][n]
            dt.append([n, familia(r), form, int(r["replica"]) if r["replica"] else None,
                       r["estado"], f(r["gap_pct"]), f(r["tiempo_scip_seg"]),
                       f(r["primal_bound"]), f(r["dual_bound"]),
                       int(float(r["nodos"])) if f(r["nodos"]) is not None else None,
                       f(r["T"]), int(float(r["suma_z"])) if f(r["suma_z"]) is not None else None,
                       r["solucion_auditada"]])
    asg = wb.create_sheet("Asignaciones")
    asg.append(["instancia", "formulacion", "estudiante", "curso_origen", "genero", "academico",
                "socioemocional", "convivencia", "curso_nuevo", "preferencias",
                "preferidos_en_mismo_curso", "satisfecho"])
    for form in ("light", "baseline"):
        for inst, r in leer_asignaciones(form, limite):
            ent = lambda v: int(v) if v.lstrip("-").isdigit() else v  # noqa: E731
            asg.append([inst, form, ent(r.get("estudiante", "")), r.get("curso_origen"),
                        r.get("genero"), ent(r.get("academico", "")), ent(r.get("socioemocional", "")),
                        ent(r.get("convivencia", "")), ent(r.get("curso_nuevo", "")),
                        r.get("preferencias"), r.get("preferidos_en_mismo_curso"),
                        ent(r.get("satisfecho", ""))])
    gris = PatternFill("solid", fgColor="DDE4EE")
    for hoja, formatos in ((ws, {6: "0.00", 7: "0.00", 8: "0.0", 9: "0.0"}),
                           (dt, {6: "0.00", 7: "0.0", 8: "0.000", 9: "0.000", 11: "0.000"}),
                           (asg, {})):
        for c in hoja[1]:
            c.font = Font(bold=True)
            c.fill = gris
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for col, fmt_ in formatos.items():
            for (celda,) in hoja.iter_rows(min_row=2, min_col=col, max_col=col):
                celda.number_format = fmt_
        for k in range(1, hoja.max_column + 1):
            hoja.column_dimensions[get_column_letter(k)].width = 22 if k == 1 else 15
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
    ws.append([])
    ws.append([f"Limite SCIP {limite} s. Desviacion estandar muestral (n-1); vacia con una sola "
               "replica. Replicas sin incumbente no entran en el gap."])
    tmp = ruta.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(ruta)   # escritura atomica: nunca queda un Excel a medio escribir
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--limite", default="3600")
    a = ap.parse_args()
    out = RAIZ / "corridas_nlhpc" / f"consolidado_T{a.limite}"
    out.mkdir(parents=True, exist_ok=True)
    datos = {form: leer(form, a.limite) for form in ("light", "baseline")}
    txt = []
    for form, filas in datos.items():
        if filas:
            escribir_csv(out / f"resultados_{form}.csv", list(filas.values()), CAMPOS)
            txt += resumen(filas, f"{form.upper()}  (limite {a.limite} s)") + [""]
        print(f"{form}: {len(filas)} instancias")
    (out / "resumen_por_familia.txt").write_text("\n".join(txt) + "\n", encoding="utf-8")
    for form in ("light", "baseline"):
        n = copiar_soluciones(form, a.limite, out / "soluciones")
        if n:
            print(f"{form}: {n} soluciones .txt copiadas a {out / 'soluciones' / form}")
    if escribir_excel(out / "resumen_resultados.xlsx", datos, a.limite):
        print(f"Excel: {out / 'resumen_resultados.xlsx'}")

    comunes = sorted(set(datos["light"]) & set(datos["baseline"]))
    if comunes:
        pares, lin = [], []
        lin += ["COMPARACION PAR A PAR (misma instancia, mismo limite)",
                "razon = tiempo baseline / tiempo light (> 1: light mas rapido); "
                "una corrida sin optimo cuenta con el limite como tiempo.", ""]
        lin.append(f"{'instancia':<22}{'est_B':>8}{'est_L':>8}{'t_B':>8}{'t_L':>8}{'razon':>8}"
                   f"{'pb_B':>8}{'pb_L':>8}{'db_B':>8}{'db_L':>8}  nota")
        for n in comunes:
            b, l = datos["baseline"][n], datos["light"][n]
            tb, tl = f(b["tiempo_scip_seg"]), f(l["tiempo_scip_seg"])
            nota = ""
            pb_b, pb_l, db_b, db_l = map(f, (b["primal_bound"], l["primal_bound"],
                                            b["dual_bound"], l["dual_bound"]))
            if b["estado"] == l["estado"] == "optimal" and abs(pb_b - pb_l) > 1e-6:
                nota = "ALERTA: optimos distintos"
            elif pb_b is not None and db_l is not None and pb_b < db_l - 1e-6:
                nota = "ALERTA: primal baseline bajo cota light"
            elif pb_l is not None and db_b is not None and pb_l < db_b - 1e-6:
                nota = "ALERTA: primal light bajo cota baseline"
            razon = tb / tl if tb and tl else None
            pares.append({"instancia": n, "familia": familia(l), "estado_baseline": b["estado"],
                          "estado_light": l["estado"], "t_baseline": tb, "t_light": tl,
                          "razon_B_L": razon, "primal_baseline": pb_b, "primal_light": pb_l,
                          "dual_baseline": db_b, "dual_light": db_l, "nota": nota})
            lin.append(f"{n:<22}{b['estado'][:7]:>8}{l['estado'][:7]:>8}{fmt(tb, 0):>8}{fmt(tl, 0):>8}"
                       f"{fmt(razon, 2):>8}{fmt(pb_b, 2):>8}{fmt(pb_l, 2):>8}{fmt(db_b, 2):>8}"
                       f"{fmt(db_l, 2):>8}  {nota}")
        rs = [p["razon_B_L"] for p in pares if p["razon_B_L"]]
        ob = sum(p["estado_baseline"] == "optimal" for p in pares)
        ol = sum(p["estado_light"] == "optimal" for p in pares)
        lin += ["", f"pares: {len(pares)} | optimas baseline: {ob} | optimas light: {ol}",
                f"mediana razon B/L: {fmt(st.median(rs), 2) if rs else '-'} | "
                f"media geometrica razon: {fmt(math.exp(st.mean(math.log(x) for x in rs)), 2) if rs else '-'}",
                f"light mas rapido (razon > 1,05) en {sum(x > 1.05 for x in rs)} de {len(rs)} pares; "
                f"baseline mas rapido (razon < 0,95) en {sum(x < 0.95 for x in rs)}"]
        escribir_csv(out / "comparacion_pares.csv", pares, list(pares[0].keys()))
        (out / "comparacion_pares.txt").write_text("\n".join(lin) + "\n", encoding="utf-8")
        print("\n".join(lin[-3:]))
    print(f"Archivos en: {out}")


if __name__ == "__main__":
    main()
